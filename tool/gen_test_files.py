"""生成 Echoic 各格式测试文件（独立脚本，不依赖项目代码）。

依赖（运行前手动安装）：
  pip install pillow python-docx pymupdf edge-tts

系统：ffmpeg（生成 test_mp4.mp4）
"""

from __future__ import annotations

import asyncio
import shutil
import subprocess
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

# 行间仅用单个换行、无空行，使 txt / pdf / docx 等后端提取结果一致。
TEXT_PLAIN = """Reading Practice
This is a test file, made for Echoic upload.
Echoic should extract this text and convert it to speech.
Well done.
Keep going every day!
Are you ready to practice?"""

OUT_DIR = Path(__file__).resolve().parent / "echoic_test_files"

# 扫描版 PDF / OCR 图片：US Letter @ 200 DPI（与 backend local.py _OCR_DPI 一致）
OCR_DPI = 200
PAGE_W, PAGE_H = int(8.5 * OCR_DPI), int(11 * OCR_DPI)
MARGIN = 120
LINE_SPACING = 16
FONT_SIZE = 42

EDGE_TTS_VOICE = "en-US-JennyNeural"
MP4_W, MP4_H = 1280, 720


def _plain_lines(text: str = TEXT_PLAIN) -> list[str]:
    return [line.strip() for line in text.splitlines() if line.strip()]


def _build_markdown() -> str:
    lines = _plain_lines()
    parts = [f"# {lines[0]}", ""]
    for line in lines[1:]:
        parts.extend([line, ""])
    return "\n".join(parts)


def _build_html() -> str:
    lines = _plain_lines()
    body = "\n".join(f"  <p>{line}</p>" for line in lines[1:])
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <title>Reading Practice Test</title>
  <style>body {{ font-family: sans-serif; }}</style>
  <script>console.log("ignored");</script>
</head>
<body>
  <h1>{lines[0]}</h1>
{body}
</body>
</html>"""


def _load_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for path in (
        Path(r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibri.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
    ):
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def _wrap_text(text: str, font: ImageFont.ImageFont, max_width: int) -> list[str]:
    lines: list[str] = []
    for paragraph in text.strip().split("\n"):
        if not paragraph.strip():
            lines.append("")
            continue
        words = paragraph.split()
        current = words[0]
        for word in words[1:]:
            trial = f"{current} {word}"
            if font.getlength(trial) <= max_width:
                current = trial
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return lines


def _render_text_image(
    text: str,
    *,
    width: int = PAGE_W,
    height: int = PAGE_H,
    font_size: int = FONT_SIZE,
    margin: int = MARGIN,
) -> Image.Image:
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font = _load_font(font_size)
    max_width = width - 2 * margin
    y = margin
    for line in _wrap_text(text, font, max_width):
        if line == "":
            y += font_size // 2
            continue
        draw.text((margin, y), line, fill="black", font=font)
        y += font_size + LINE_SPACING
    return img


def _render_mp4_frame(text: str) -> Image.Image:
    """与 OCR 测试图同源渲染，等比缩放至 1280×720。"""
    source = _render_text_image(text)
    canvas = Image.new("RGB", (MP4_W, MP4_H), "white")
    scale = min(MP4_W / source.width, MP4_H / source.height)
    size = (int(source.width * scale), int(source.height * scale))
    resized = source.resize(size, Image.LANCZOS)
    offset = ((MP4_W - size[0]) // 2, (MP4_H - size[1]) // 2)
    canvas.paste(resized, offset)
    return canvas


def _clean_out_dir() -> None:
    if OUT_DIR.exists():
        for item in OUT_DIR.iterdir():
            if item.is_file():
                item.unlink()
            elif item.is_dir():
                shutil.rmtree(item)


def write_test_txt() -> Path:
    path = OUT_DIR / "test_txt.txt"
    path.write_text(TEXT_PLAIN, encoding="utf-8")
    return path


def write_test_md() -> Path:
    path = OUT_DIR / "test_md.md"
    path.write_text(_build_markdown(), encoding="utf-8")
    return path


def write_test_html() -> Path:
    path = OUT_DIR / "test_html.html"
    path.write_text(_build_html(), encoding="utf-8")
    return path


def write_test_htm() -> Path:
    path = OUT_DIR / "test.htm"
    path.write_text(_build_html(), encoding="utf-8")
    return path


def write_test_docx() -> Path:
    from docx import Document

    lines = _plain_lines()
    path = OUT_DIR / "test_docx.docx"
    doc = Document()
    doc.add_heading(lines[0], level=1)
    for line in lines[1:]:
        doc.add_paragraph(line)
    doc.save(path)
    return path


def write_test_pdf() -> Path:
    """文字版 PDF（有文字层）"""
    import fitz

    path = OUT_DIR / "test_pdf.pdf"
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    y = 72
    for line in _plain_lines():
        page.insert_text((72, y), line, fontsize=14, fontname="helv")
        y += 22
    doc.save(path)
    doc.close()
    return path


def write_test_pdf_scanned() -> Path:
    """扫描版 PDF（纯图片，无文字层）"""
    path = OUT_DIR / "test_pdf_scanned.pdf"
    img = _render_text_image(TEXT_PLAIN)
    img.save(path, "PDF", resolution=float(OCR_DPI))
    return path


def write_test_png() -> Path:
    path = OUT_DIR / "test_png.png"
    img = _render_text_image(TEXT_PLAIN)
    img.save(path, "PNG", dpi=(OCR_DPI, OCR_DPI))
    return path


def write_test_jpg() -> Path:
    path = OUT_DIR / "test_jpg.jpg"
    img = _render_text_image(TEXT_PLAIN)
    img.save(path, "JPEG", quality=92, dpi=(OCR_DPI, OCR_DPI))
    return path


def _tts_text(text: str) -> str:
    """与 TEXT_PLAIN 一致；仅给无标点的行补句号，便于 TTS 断句。"""
    parts: list[str] = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if line[-1] not in ".!?":
            line += "."
        parts.append(line)
    return " ".join(parts)


async def _edge_tts_save(text: str, output_path: Path) -> None:
    import edge_tts

    communicate = edge_tts.Communicate(text, EDGE_TTS_VOICE)
    await communicate.save(str(output_path))


def _synthesize_tts_audio(text: str, audio_path: Path) -> None:
    try:
        asyncio.run(_edge_tts_save(_tts_text(text), audio_path))
    except ImportError as e:
        raise RuntimeError("edge-tts not installed") from e
    except Exception as e:
        raise RuntimeError(f"edge-tts failed: {e}") from e
    if not audio_path.exists() or audio_path.stat().st_size == 0:
        raise RuntimeError("edge-tts produced empty audio file")


def write_test_mp3() -> Path:
    path = OUT_DIR / "test_mp3.mp3"
    _synthesize_tts_audio(TEXT_PLAIN, path)
    return path


def write_test_mp4() -> Path:
    """视频 MP4（静态画面 + edge-tts 语音，测试 ffmpeg 提音轨）"""
    path = OUT_DIR / "test_mp4.mp4"
    frame_path = OUT_DIR / "_tmp_mp4_frame.png"
    audio_path = OUT_DIR / "test_mp3.mp3"

    if not audio_path.exists():
        raise RuntimeError("test_mp3.mp3 not found")

    _render_mp4_frame(TEXT_PLAIN).save(frame_path)

    try:
        subprocess.run(
            [
                "ffmpeg", "-y",
                "-loop", "1", "-i", str(frame_path),
                "-i", str(audio_path),
                "-c:v", "libx264", "-tune", "stillimage",
                "-c:a", "aac", "-b:a", "128k",
                "-pix_fmt", "yuv420p",
                "-movflags", "+faststart",
                "-shortest",
                str(path),
            ],
            check=True,
            capture_output=True,
        )
    except FileNotFoundError as e:
        raise RuntimeError("ffmpeg not found") from e
    except subprocess.CalledProcessError as e:
        stderr = (e.stderr or b"").decode(errors="replace").strip()
        raise RuntimeError(f"ffmpeg failed: {stderr}") from e
    finally:
        frame_path.unlink(missing_ok=True)

    return path


def main() -> None:
    _clean_out_dir()
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    writers = [
        ("TXT", write_test_txt),
        ("Markdown", write_test_md),
        ("HTML", write_test_html),
        ("HTM", write_test_htm),
        ("DOCX", write_test_docx),
        ("PDF (text layer)", write_test_pdf),
        ("PDF (scanned/image)", write_test_pdf_scanned),
        ("PNG (OCR)", write_test_png),
        ("JPEG (OCR)", write_test_jpg),
        ("MP3 (audio)", write_test_mp3),
        ("MP4 (video+audio)", write_test_mp4),
    ]

    print(f"Output directory: {OUT_DIR.resolve()}\n")
    failures = 0
    for label, fn in writers:
        try:
            path = fn()
            print(f"[OK] {label:<22} -> {path.name}")
        except Exception as e:
            failures += 1
            print(f"[FAIL] {label:<22} -> {e}")

    if failures:
        sys.exit(1)


if __name__ == "__main__":
    main()
